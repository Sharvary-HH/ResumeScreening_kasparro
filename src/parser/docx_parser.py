"""DOCX parsing.

The corpus is all PDF, but resumes arrive as .docx often enough that the
pipeline should handle both. Same trap as PDFs: hyperlink targets are not in the
paragraph text, they live in the document relationships.
"""
from pathlib import Path

import docx

from src.models import ParsedResume
from src.parser.urls import extract_urls_from_text, merge_urls


def parse_docx(path: Path) -> ParsedResume:
    path = Path(path)
    try:
        document = docx.Document(str(path))

        blocks = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                blocks.extend(cell.text for cell in row.cells)

        rel_urls = [
            rel.target_ref
            for rel in document.part.rels.values()
            if rel.reltype.endswith("/hyperlink") and rel.target_ref
        ]
    except Exception as exc:  # noqa: BLE001
        return ParsedResume(
            source_path=str(path),
            parser="docx",
            error=f"DOCX read failed: {exc}",
        )

    text = "\n".join(block for block in blocks if block and block.strip())
    return ParsedResume(
        source_path=str(path),
        text=text,
        urls=merge_urls(rel_urls, extract_urls_from_text(text)),
        page_count=1,  # DOCX has no fixed pagination without rendering it
        parser="docx",
    )
