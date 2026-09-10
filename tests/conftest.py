from pathlib import Path

import pytest

FIXTURES = Path(__file__).parent / "fixtures"
RESUMES = Path(__file__).parent.parent / "resumes"


@pytest.fixture(scope="session")
def synthetic_docx() -> Path:
    """Build a small .docx once per session.

    The corpus is PDF-only, so the DOCX path needs a fixture of its own. It is
    generated rather than committed so the test stays readable - you can see
    exactly what it contains right here.
    """
    import docx

    FIXTURES.mkdir(parents=True, exist_ok=True)
    path = FIXTURES / "synthetic_resume.docx"

    document = docx.Document()
    document.add_heading("Priya Nair", level=1)
    document.add_paragraph("priya.nair@example.com | Bengaluru, India")
    document.add_paragraph("GitHub")  # the URL lives in the relationship, not here
    document.add_heading("Skills", level=2)
    document.add_paragraph("Python, FastAPI, LangGraph, PostgreSQL, Docker")
    document.add_heading("Projects", level=2)
    document.add_paragraph(
        "Support Copilot - a multi-agent RAG assistant over internal docs, "
        "built with LangGraph and FAISS behind an async FastAPI service."
    )

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Internship"
    table.cell(0, 1).text = "Backend intern at Acme, built Redis-backed job queues."

    # python-docx only writes a hyperlink relationship via the part API.
    document.part.relate_to(
        "https://github.com/priya-nair",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )

    document.save(str(path))
    return path
